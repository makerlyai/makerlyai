from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

# --- Page margins ---
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

def h1(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x00, 0x00, 0x00)
    return p

def h2(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x1a, 0x4b, 0x9c)
    return p

def h3(text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
    return p

def body(text, italic=False):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.italic = italic
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    return p

def bullet(text):
    p = doc.add_paragraph(style='List Bullet')
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    return p

def hr():
    doc.add_paragraph("─" * 70)

def spacer():
    doc.add_paragraph("")

# ══════════════════════════════════════════════════════════
# PAGE 1 — Cover
# ══════════════════════════════════════════════════════════
p = doc.add_paragraph()
run = p.add_run("makerlyai")
run.bold = True
run.font.size = Pt(36)
run.font.color.rgb = RGBColor(0x1a, 0x4b, 0x9c)

body("Portfolio & Rate Card 2026")
body("D I G I T A L  A R C H I T E C T U R E  &  A I  E N G I N E E R I N G", italic=True)

spacer()
h1("We build what you grow.")

spacer()
body("Makerly AI is a software development agency for founders and businesses that need SaaS products, AI agents, web apps, mobile apps, and automation systems built with speed and clarity.")

spacer()
stats = doc.add_paragraph()
stats.add_run("⚡  24h  ").bold = True
stats.add_run("First working preview     ")
stats.add_run("5+  ").bold = True
stats.add_run("Products shipped     ")
stats.add_run("100%  ").bold = True
stats.add_run("Satisfaction guarantee")

spacer()
body("EMAIL    getmakerlyai@gmail.com")
body("WEB      makerlyai.in")
body("LOCATION  Jamshedpur, Jharkhand, India")

doc.add_page_break()

# ══════════════════════════════════════════════════════════
# PAGE 2 — Capabilities
# ══════════════════════════════════════════════════════════
h2("CAPABILITIES")
h2("W H A T  W E  B U I L D")
h3("Premium Deliverables")
body("We don't build standard software. Every product we deploy is engineered for speed, conversion, and undeniable aesthetic superiority.")

spacer()
h2("O U R  P R O C E S S")
h3("From Idea to Product")

steps = [
    ("01", "Tell Us Your Idea", "Share your vision — app, SaaS, AI agent, website. We listen, ask sharp questions, and map your requirements."),
    ("02", "We Build in 24 Hours", "Our elite engineers deliver a working preview within 24 hours. Real code. Real product. Not a mockup."),
    ("03", "Pay Only If You Like It", "Review the preview. If it doesn't meet your standards, you pay nothing. Zero risk, maximum reward."),
]
for num, title, desc in steps:
    p = doc.add_paragraph()
    run = p.add_run(f"{num}  {title}")
    run.bold = True
    run.font.size = Pt(12)
    run.font.color.rgb = RGBColor(0x1a, 0x4b, 0x9c)
    body(desc)
    spacer()

doc.add_page_break()

# ══════════════════════════════════════════════════════════
# PAGE 3 — What We Build (Services)
# ══════════════════════════════════════════════════════════
services = [
    ("🧠 AI Agents & Automation", "Deploy intelligent virtual employees that handle support, scheduling, and sales 24/7. Custom LLM integration with OpenAI, Anthropic, and Groq."),
    ("🌐 Immersive Web Experiences", "Blazing fast marketing sites that feel like native apps — fluid, interactive, conversion-optimized. Built with Next.js, React, and modern frameworks."),
    ("📱 Cross-Platform Apps", "Native-quality mobile and desktop apps built with modern frameworks. One codebase, every device. Android, iOS, and Web."),
    ("📈 SaaS & Revenue Engines", "Scalable multi-tenant platforms with billing, analytics, and growth tools built in from day one. Architected for MRR."),
]
for title, desc in services:
    h3(title)
    body(desc)
    spacer()

doc.add_page_break()

# ══════════════════════════════════════════════════════════
# PAGE 4 — INVESTMENT & RATE CARD
# ══════════════════════════════════════════════════════════
h2("I N V E S T M E N T  &  E N G A G E M E N T  M O D E L S")
h3("Rate Card")
body("Flexible engagement models tailored for maximum ROI. Whether you need ongoing leadership, dedicated sprints, or an end-to-end build — we adapt.")

spacer()
h2("P R O J E C T - B A S E D  P A C K A G E S")

# --- LAUNCH PACKAGE ---
hr()
p = doc.add_paragraph()
r = p.add_run("🚀  Launch Package  —  INR 25,000 starting")
r.bold = True
r.font.size = Pt(14)
r.font.color.rgb = RGBColor(0x1a, 0x4b, 0x9c)

body("Ideal for landing pages, MVPs, and rapid product launches. Get to market fast.")

table_launch = doc.add_table(rows=2, cols=2)
table_launch.style = 'Table Grid'
table_launch.cell(0, 0).text = "DEVELOPER"
table_launch.cell(0, 1).text = "UI/UX DESIGN"
table_launch.cell(1, 0).text = "₹1,500 /hr"
table_launch.cell(1, 1).text = "₹2,000 /hr"

spacer()
features_launch = [
    "Up to 10 pages / screens",
    "High-fidelity responsive UI/UX",
    "Database integration (Supabase / Prisma)",
    "Standard API & form setup",
    "SEO & performance optimization",
    "Delivery in ~1 week",
]
for f in features_launch:
    bullet(f)

spacer()

# --- SCALE PACKAGE ---
hr()
p = doc.add_paragraph()
r = p.add_run("⭐  Scale Package  —  Recommended  —  INR 50,000+ starting")
r.bold = True
r.font.size = Pt(14)
r.font.color.rgb = RGBColor(0x1a, 0x4b, 0x9c)

body("For comprehensive platforms, complex web-apps, and full AI integrations.")

table_scale = doc.add_table(rows=2, cols=2)
table_scale.style = 'Table Grid'
table_scale.cell(0, 0).text = "DEVELOPER"
table_scale.cell(0, 1).text = "UI/UX DESIGN"
table_scale.cell(1, 0).text = "₹4,000 /hr"
table_scale.cell(1, 1).text = "₹2,800 /hr"

spacer()
features_scale = [
    "Unlimited core pages",
    "Custom AI workflows & LLM integrations",
    "Authentication & payment gateways",
    "Dynamic dashboards & CMS",
    "Complete CI/CD & cloud deployment",
    "Delivery in 3–5 weeks",
]
for f in features_scale:
    bullet(f)

doc.add_page_break()

# ══════════════════════════════════════════════════════════
# PAGE 5 — Tech Stack & Contact
# ══════════════════════════════════════════════════════════
h2("T E C H N O L O G Y  S T A C K")
h3("Built With the Best")

stacks = {
    "Frontend": ["Next.js 14 / React 18", "TypeScript", "Tailwind CSS", "Framer Motion"],
    "Backend & Data": ["Node.js / Express", "Prisma ORM", "Supabase / PostgreSQL", "REST & GraphQL APIs"],
    "AI & Cloud": ["OpenAI / Anthropic / Groq", "Custom RAG Pipelines", "Vercel / AWS", "CI/CD Automation"],
}
for category, items in stacks.items():
    h3(category)
    for item in items:
        bullet(item)
    spacer()

spacer()
h2("W H Y  M A K E R L Y  A I")
why = [
    ("🔥 Speed", "Working preview in 24 hours. No endless discovery phases."),
    ("💎 Premium Quality", "Every build feels expensive, runs fast, and forces competition to play catch-up."),
    ("🛡️ Zero Risk", "Pay only if you like the result. Your satisfaction is guaranteed."),
]
for icon_title, desc in why:
    h3(icon_title)
    body(desc)
    spacer()

hr()
h1("Ready to build your digital empire?")
body("Let's start the conversation. Share your vision and we'll deliver the first working preview within 24 hours.")
spacer()
body("EMAIL    getmakerlyai@gmail.com")
body("WEB      makerlyai.in")
body("Social:  Instagram  |  X (Twitter)  |  LinkedIn  |  GitHub  |  YouTube  |  Telegram")
spacer()
body("MAKERLY AI © 2026 — CONFIDENTIAL & PROPRIETARY", italic=True)

# ══════════════════════════════════════════════════════════
# SAVE
# ══════════════════════════════════════════════════════════
output_path = r'C:\Users\windows\Desktop\Makerly_AI_Portfolio_and_Rates.docx'
doc.save(output_path)
print(f"Word document saved to: {output_path}")
